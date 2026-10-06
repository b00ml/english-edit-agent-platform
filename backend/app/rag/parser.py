"""Structure-aware parsing. parse_file remains the legacy string interface."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Set
from xml.etree import ElementTree

from app.errors import DocumentParseError as DocumentParseError
from app.rag import tables
from app.rag.document import (
    BlockType,
    DocumentBlock,
    ParsedDocument,
    block_text,
    finalize_document,
)

PARSER_VERSION = "rag-parser-v2"
SUPPORTED_EXTENSIONS: Set[str] = {
    ".txt",
    ".md",
    ".markdown",
    ".docx",
    ".pdf",
    ".html",
    ".htm",
    ".csv",
    ".xlsx",
}
MAX_EXPANDED_BYTES = 40 * 1024 * 1024
MAX_BLOCKS = 50000
MAX_CELLS = 1000000


class UnsupportedFileTypeError(ValueError):
    """The input extension is not supported."""


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("gb18030", errors="replace")


def _block(
    kind: BlockType,
    text: str,
    *,
    level: int | None = None,
    page: int | None = None,
    locator: dict[str, Any] | None = None,
    meta: dict[str, Any] | None = None,
) -> DocumentBlock:
    return {
        "block_id": "",
        "block_type": kind,
        "text": text,
        "level": level,
        "page_no": page,
        "order": 0,
        "source_locator": locator or {},
        "meta": meta or {},
    }


def _finish(
    filename: str,
    blocks: list[DocumentBlock],
    warnings: list[str],
    *,
    original_text: str | None = None,
    stats: dict[str, Any] | None = None,
) -> ParsedDocument:
    if len(blocks) > MAX_BLOCKS:
        raise DocumentParseError("解析 block 数超过安全上限")
    for index, block in enumerate(blocks):
        block["order"] = index
        block["block_id"] = f"block:{index}"
    return finalize_document(
        {
            "source_name": Path(filename).stem,
            "source_type": Path(filename).suffix.lower()[1:],
            "parser_version": PARSER_VERSION,
            "blocks": blocks,
            "warnings": warnings,
            "stats": stats or {},
            "text": "",
            "source_hash": "",
            "original_text": original_text,
        }
    )


def parse_file(filename: str, content: bytes) -> str:
    return block_text(parse_document(filename, content))


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(f"不支持的文件类型 {ext or '（无扩展名）'}")
    if len(content) > 10 * 1024 * 1024:
        raise DocumentParseError("文件超过 10MB 安全上限")
    if ext in {".docx", ".xlsx"}:
        _check_archive(content)
    try:
        if ext in {".txt", ".md", ".markdown"}:
            doc = parse_text(filename, _decode_text(content))
        elif ext == ".docx":
            doc = _parse_docx(filename, content)
        elif ext == ".pdf":
            doc = _parse_pdf(filename, content)
        elif ext in {".html", ".htm"}:
            doc = _parse_html(filename, _decode_text(content))
        elif ext == ".csv":
            text = _decode_text(content)
            rows = [(index, row) for index, row in enumerate(csv.reader(io.StringIO(text)), 1)]
            doc = _finish(
                filename, _sheet_blocks(Path(filename).stem, rows), [], original_text=text
            )
        else:
            doc = _parse_xlsx(filename, content)
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError, csv.Error) as exc:
        raise DocumentParseError(f"文件解析失败: {type(exc).__name__}") from exc
    if ext in {".txt", ".md", ".markdown", ".html", ".htm", ".csv"}:
        try:
            content.decode("utf-8-sig")
        except UnicodeDecodeError:
            doc["warnings"].append("文件编码回退到 GB18030，无法识别的字符可能被替换，请核对原文")
    doc["source_hash"] = hashlib.sha256(content).hexdigest()
    return doc


def _check_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 4096
                or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES
            ):
                raise DocumentParseError("文档解压大小或文件数超过安全上限")
    except zipfile.BadZipFile as exc:
        raise DocumentParseError("文档不是有效的 Office ZIP 文件") from exc


def parse_text(filename: str, text: str) -> ParsedDocument:
    """Identify Markdown headings/tables/fences, preserving raw text snapshots."""
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = normalized.splitlines()
    blocks: list[DocumentBlock] = []
    warnings: list[str] = []
    index = 0
    while index < len(lines):
        if not lines[index].strip():
            index += 1
            continue
        start = index
        heading = re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$", lines[index])
        fence = re.match(r"^\s*(`{3,}|~{3,})", lines[index])
        if heading:
            blocks.append(
                _block(
                    "heading",
                    lines[index],
                    level=len(heading[1]),
                    locator={"line_start": index + 1, "line_end": index + 1},
                    meta={"title": heading[2]},
                )
            )
            index += 1
        elif tables.is_table_start(lines, index):
            header = tables.cells(lines[index])
            width = len(header)
            index += 2
            rows = [header]
            while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                if tables.is_table_start(lines, index) or len(tables.cells(lines[index])) != width:
                    break
                rows.append(tables.cells(lines[index]))
                index += 1
            value = tables.render(rows, already_escaped=True)
            blocks.append(
                _block(
                    "table",
                    value,
                    locator={"line_start": start + 1, "line_end": index},
                    meta={
                        "table_id": f"md-table:{start + 1}",
                        "rows": rows,
                        "row_count": len(rows) - 1,
                        "column_count": width,
                        "original_text": "\n".join(lines[start:index]),
                    },
                )
            )
        elif fence or lines[index].strip() == "$$":
            marker = fence[1] if fence else "$$"
            index += 1
            while index < len(lines) and not lines[index].strip().startswith(marker):
                index += 1
            closed = index < len(lines)
            if closed:
                index += 1
            else:
                warnings.append(f"第 {start + 1} 行开始的代码/公式没有闭合")
            blocks.append(
                _block(
                    "code",
                    "\n".join(lines[start:index]),
                    locator={"line_start": start + 1, "line_end": index},
                    meta={"fence": marker, "closed": closed},
                )
            )
        else:
            index += 1
            while index < len(lines) and lines[index].strip():
                if (
                    re.match(r"^\s{0,3}#{1,6}\s", lines[index])
                    or tables.is_table_start(lines, index)
                    or re.match(r"^\s*(`{3,}|~{3,}|\$\$)", lines[index])
                ):
                    break
                index += 1
            value = "\n".join(lines[start:index])
            kind: BlockType = "list" if re.match(r"^\s*(?:[-*+] |\d+[.)] )", value) else "paragraph"
            blocks.append(_block(kind, value, locator={"line_start": start + 1, "line_end": index}))
    return _finish(filename, blocks, warnings, original_text=text)


def _parse_docx(filename: str, content: bytes) -> ParsedDocument:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from lxml.etree import XMLSyntaxError

    try:
        document = Document(io.BytesIO(content))
    except (XMLSyntaxError, PackageNotFoundError, ValueError) as exc:
        raise DocumentParseError(f"DOCX 解析失败: {type(exc).__name__}") from exc
    blocks: list[DocumentBlock] = []
    warnings: list[str] = []
    table_index = 0
    body = document.element.body
    if any(
        part.partname.endswith(("header1.xml", "footer1.xml"))
        for part in document.part.package.parts
    ):
        warnings.append("DOCX 页眉/页脚未索引，保留正文；请核对页眉页脚中的重要信息")
    if body.xpath(".//w:drawing | .//w:pict"):
        warnings.append("DOCX 图片未执行 OCR；图片文字不能作为已索引文本")
    if body.xpath(".//w:txbxContent"):
        warnings.append("DOCX 文本框布局未还原，请核对原文")
    if body.xpath(".//m:oMath"):
        warnings.append("DOCX Office 公式未结构化提取，请核对原文")
    for body_order, child in enumerate(body.iterchildren()):
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            if not paragraph.text.strip():
                continue
            style = paragraph.style.name if paragraph.style else ""
            level_match = re.search(r"(?:Heading|标题)\s*([1-6])", style, re.I)
            level = int(level_match[1]) if level_match else None
            blocks.append(
                _block(
                    "heading" if level else "paragraph",
                    paragraph.text,
                    level=level,
                    locator={"body_order": body_order},
                    meta={"style": style, "title": paragraph.text if level else None},
                )
            )
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows = []
            for row in table.rows:
                values = [""] * getattr(row, "grid_cols_before", 0)
                for cell in row.cells:
                    value = _docx_cell_text(cell, warnings, table_index)
                    values.append(value)
                values.extend([""] * getattr(row, "grid_cols_after", 0))
                rows.append(values)
            inferred_header = not bool(child.xpath("./w:tr[1]/w:trPr/w:tblHeader"))
            if inferred_header:
                warnings.append(
                    f"DOCX 表 {table_index} 首行暂按表头展示，原文件未声明重复表头，请核对"
                )
            merged = bool(child.xpath(".//w:gridSpan | .//w:vMerge"))
            if merged:
                warnings.append(f"DOCX 表 {table_index} 有合并单元格，按网格重复内容；非视觉原版")
            blocks.append(
                _block(
                    "table",
                    tables.render(rows),
                    locator={"body_order": body_order, "table_index": table_index},
                    meta={
                        "table_id": f"docx-table:{table_index}",
                        "table_index": table_index,
                        "rows": rows,
                        "row_count": max(0, len(rows) - 1),
                        "column_count": max((len(row) for row in rows), default=0),
                        "merged_cells": merged,
                        "header_inferred": inferred_header,
                        "cells": [
                            {"row_index": r, "column_index": c, "text": value}
                            for r, row in enumerate(rows)
                            for c, value in enumerate(row)
                        ],
                        "original_xml": child.xml,
                    },
                )
            )
            table_index += 1
    return _finish(filename, blocks, list(dict.fromkeys(warnings)))


def _docx_cell_text(cell: Any, warnings: list[str], table_index: int, depth: int = 0) -> str:
    """Preserve nested cell text; refuse pathological depth instead of dropping it."""
    if depth > 8:
        raise DocumentParseError("DOCX 嵌套表格超过安全深度")
    parts = [cell.text]
    if cell.tables:
        warnings.append(f"DOCX 表 {table_index} 含嵌套表格，已保留文本但布局降级")
        for nested in cell.tables:
            for row in nested.rows:
                parts.append(
                    " / ".join(
                        _docx_cell_text(child, warnings, table_index, depth + 1)
                        for child in row.cells
                    )
                )
    return "\n".join(parts)


def _parse_pdf(filename: str, content: bytes) -> ParsedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError, PdfStreamError

    try:
        reader = PdfReader(io.BytesIO(content))
        pages = reader.pages
        page_count = len(pages)
    except (PdfReadError, PdfStreamError, ValueError) as exc:
        raise DocumentParseError(f"PDF 解析失败: {type(exc).__name__}") from exc
    blocks: list[DocumentBlock] = []
    warnings = ["PDF 使用普通文本抽取，未承诺双栏、表格布局或 OCR 正确性"]
    for page_no, page in enumerate(pages, 1):
        try:
            text = page.extract_text() or ""
        except (PdfReadError, PdfStreamError, ValueError, TypeError) as exc:
            text = ""
            warnings.append(f"PDF 第 {page_no} 页抽取失败: {type(exc).__name__}")
        if not text.strip():
            warnings.append(f"PDF 第 {page_no} 页没有可抽取文本，可能为扫描/空页；未执行 OCR")
            blocks.append(_block("page_break", "", page=page_no, locator={"page_no": page_no}))
        else:
            blocks.append(
                _block(
                    "paragraph",
                    text,
                    page=page_no,
                    locator={"page_no": page_no},
                    meta={"parser_quality": "text"},
                )
            )
    return _finish(filename, blocks, warnings, stats={"page_count": page_count})


class _HTMLCollector(HTMLParser):
    """Small HTML-to-block adapter; merged cells retain span metadata and warnings."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[DocumentBlock] = []
        self.warnings: list[str] = []
        self.buffer: list[str] = []
        self.heading: int | None = None
        self.skip_depth = 0
        self.table_depth = 0
        self.rows: list[list[str]] = []
        self.row: list[str] = []
        self.cell: list[str] | None = None
        self.spans: list[dict[str, str | None]] = []

    def flush(self) -> None:
        text = re.sub(r"\s+", " ", "".join(self.buffer)).strip()
        if text:
            self.blocks.append(
                _block(
                    "heading" if self.heading else "paragraph",
                    text,
                    level=self.heading,
                    meta={"title": text},
                )
            )
        self.buffer = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self.skip_depth += 1
            return
        if self.skip_depth:
            return
        if tag == "img":
            self.warnings.append("HTML 图片未执行 OCR")
        if tag == "table":
            if self.table_depth:
                self.warnings.append("HTML 嵌套表格已降级，需人工核对布局")
            else:
                self.flush()
                self.rows = []
                self.spans = []
            self.table_depth += 1
        elif self.table_depth and tag == "tr":
            if self.table_depth == 1:
                self.row = []
        elif self.table_depth and tag in {"td", "th"}:
            if self.table_depth == 1:
                self.cell = []
                attributes = dict(attrs)
                if attributes.get("rowspan") or attributes.get("colspan"):
                    self.spans.append(attributes)
        elif not self.table_depth:
            if re.fullmatch(r"h[1-6]", tag):
                self.flush()
                self.heading = int(tag[1])
            elif tag in {"p", "div", "li", "br"}:
                self.flush()

    def handle_data(self, data: str) -> None:
        if self.skip_depth:
            return
        if self.table_depth and self.cell is not None:
            self.cell.append(data)
        elif not self.table_depth:
            self.buffer.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self.skip_depth = max(0, self.skip_depth - 1)
            return
        if self.skip_depth:
            return
        if self.table_depth == 1 and tag in {"td", "th"} and self.cell is not None:
            self.row.append(" ".join("".join(self.cell).split()))
            self.cell = None
        elif self.table_depth == 1 and tag == "tr":
            self.rows.append(self.row)
            self.row = []
        elif tag == "table" and self.table_depth:
            self.table_depth -= 1
            if not self.table_depth:
                if self.spans:
                    self.warnings.append(
                        "HTML 合并单元格未展开，保留原始 HTML 和 rowspan/colspan metadata"
                    )
                self.blocks.append(
                    _block(
                        "table",
                        tables.render(self.rows),
                        meta={
                            "table_id": f"html-table:{len(self.blocks)}",
                            "rows": self.rows,
                            "spans": self.spans,
                            "row_count": max(0, len(self.rows) - 1),
                        },
                    )
                )
        elif not self.table_depth:
            if re.fullmatch(r"h[1-6]", tag):
                self.flush()
                self.heading = None
            elif tag in {"p", "div", "li"}:
                self.flush()


def _parse_html(filename: str, text: str) -> ParsedDocument:
    collector = _HTMLCollector()
    collector.feed(text)
    collector.close()
    collector.flush()
    if collector.table_depth:
        collector.warnings.append("HTML 表格没有闭合，无法完整还原")
        if collector.cell:
            collector.row.append("".join(collector.cell))
        if collector.row:
            collector.rows.append(collector.row)
        collector.blocks.append(_block("table", tables.render(collector.rows)))
    return _finish(filename, collector.blocks, collector.warnings, original_text=text)


def _sheet_blocks(name: str, rows: list[tuple[int, list[str]]]) -> list[DocumentBlock]:
    rows = [(index, row) for index, row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return []
    header_index, headers = rows[0]
    width = max(len(row) for _, row in rows)
    headers = headers + [""] * (width - len(headers))
    labels = [value or f"列{index + 1}" for index, value in enumerate(headers)]
    blocks = [
        _block(
            "paragraph",
            f"Sheet：{name}\n列：" + "、".join(labels),
            locator={"sheet": name, "row": header_index},
            meta={"chunk_type": "table_summary", "headers": labels},
        )
    ]
    blocks.append(
        _block(
            "paragraph",
            "列说明：\n"
            + "\n".join(f"第{index + 1}列：{label}" for index, label in enumerate(labels)),
            locator={"sheet": name, "row": header_index},
            meta={"chunk_type": "table_column", "headers": labels},
        )
    )
    for row_index, row in rows[1:]:
        if len(blocks) >= MAX_BLOCKS:
            raise DocumentParseError("表格行数超过安全上限")
        values = row + [""] * (width - len(row))
        blocks.append(
            _block(
                "table_row",
                f"Sheet：{name}\n行号：{row_index}\n"
                + "\n".join(f"{label}：{value}" for label, value in zip(labels, values)),
                locator={"sheet": name, "row": row_index},
                meta={"chunk_type": "table_row", "headers": labels, "values": values},
            )
        )
    return blocks


def _parse_xlsx(filename: str, content: bytes) -> ParsedDocument:
    from datetime import date, datetime

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    except (InvalidFileException, ValueError, ElementTree.ParseError) as exc:
        raise DocumentParseError(f"XLSX 解析失败: {type(exc).__name__}") from exc
    blocks: list[DocumentBlock] = []
    warnings: list[str] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row * sheet.max_column > MAX_CELLS:
                raise DocumentParseError("XLSX worksheet 网格超过安全上限")
            if sheet.merged_cells.ranges:
                warnings.append(
                    f"XLSX Sheet {sheet.title} 含合并单元格，保留原值和范围，不推测空位值"
                )
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                values = []
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        warnings.append(
                            f"XLSX {sheet.title}!{cell.coordinate} 保留公式文本，未计算公式结果"
                        )
                    values.append(
                        value.isoformat()
                        if isinstance(value, (date, datetime))
                        else "" if value is None else str(value)
                    )
                rows.append((row_index, values))
            sheet_blocks = _sheet_blocks(sheet.title, rows)
            for block in sheet_blocks:
                block["meta"]["merged_ranges"] = [str(value) for value in sheet.merged_cells.ranges]
            if not sheet_blocks:
                warnings.append(f"XLSX Sheet {sheet.title} 为空")
            blocks.extend(sheet_blocks)
    finally:
        workbook.close()
    return _finish(filename, blocks, warnings, stats={"sheet_count": len(workbook.worksheets)})
