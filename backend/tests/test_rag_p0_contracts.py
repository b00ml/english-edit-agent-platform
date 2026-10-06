"""No-network production-path regressions for RAG P0 A/B/C."""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from sqlalchemy.exc import SQLAlchemyError

from app.config import settings
from app.errors import (
    ChunkingError,
    DocumentParseError,
    InvalidKnowledgeError,
    TenantScopeDeniedError,
)
from app.models import KnowledgeChunk, KnowledgeDocument
from app.rag import retriever
from app.rag.chunking import Chunk, ChunkingConfig, split_document, split_text
from app.rag.chunking.validator import coverage_ratio, validate_chunks
from app.rag.indexer import index_document
from app.rag.parser import parse_document, parse_text
from app.schemas import KnowledgeChunkOut, KnowledgeUploadOut
from app.services.knowledge_service import KnowledgeService


@pytest.fixture(autouse=True)
def legacy_contracts(monkeypatch):
    # P0 tests intentionally assert contiguous ranges and legacy snippet compatibility.
    monkeypatch.setattr(settings, "RAG_CHUNK_LAYOUT", "legacy")
    monkeypatch.setattr(settings, "RAG_CONTEXT_MODE", "legacy")


def researcher(tenant=None):
    return SimpleNamespace(id="user", role="researcher", tenant_id=tenant, status="active")


def vectors(inputs, **kwargs):
    return [[1.0] * 1024 for _ in inputs]


def assert_contract(doc, chunks, cfg):
    assert validate_chunks(chunks, doc["text"], cfg.normalized()) == []
    assert coverage_ratio(chunks, doc["text"]) == 1.0
    for chunk in chunks:
        assert doc["text"][chunk.start : chunk.end] == chunk.content
        assert len(chunk.embedding_content) <= cfg.chunk_size


@pytest.mark.parametrize("strategy", ["auto", "heading", "heuristic", "recursive", "legacy"])
def test_strategies_cover_every_source_character(strategy):
    doc = parse_text("grammar.md", "# Chapter\n\n" + "Rule one. Rule two! 中文说明。\n" * 40)
    cfg = ChunkingConfig(strategy=strategy, chunk_size=170, chunk_overlap=25, min_chunk_chars=20)
    chunks, diag = split_document(doc, cfg)
    assert_contract(doc, chunks, cfg)
    assert diag.chunk_count == len(chunks)


def test_validated_fallback(monkeypatch):
    from app.rag.chunking import strategy

    original = strategy._run_tier

    def broken(doc, cfg, tier):
        return [] if tier == "heading" else original(doc, cfg, tier)

    monkeypatch.setattr(strategy, "_run_tier", broken)
    chunks, diag = split_text("# Chapter\n\nImportant rule")
    assert chunks and diag.strategy_used == "heuristic"
    assert "coverage" in diag.fallback_attempts[0]["reason"]


def test_all_strategies_fail_explicitly(monkeypatch):
    from app.rag.chunking import strategy

    monkeypatch.setattr(strategy, "_run_tier", lambda *args: [])
    with pytest.raises(ChunkingError):
        split_text("important rule")


def test_validator_uses_union_of_source_ranges_not_duplicate_bytes():
    chunks = [Chunk(content="AAAA", start=0, end=4), Chunk(content="AAAA", start=0, end=4)]
    assert coverage_ratio(chunks, "AAAA BBBB") == 0.5
    assert "incomplete_non_whitespace_coverage" in validate_chunks(
        chunks, "AAAA BBBB", ChunkingConfig()
    )


@pytest.mark.parametrize("text", ["A" * 480 + "\n" + "B" * 900, "word " * 1000, "规则。" * 500])
def test_pathological_boundaries_progress(text):
    doc = parse_text("rules.txt", text)
    cfg = ChunkingConfig(chunk_size=500, chunk_overlap=50, min_chunk_chars=80)
    chunks, _ = split_document(doc, cfg)
    assert len(chunks) < len(doc["text"]) / 100 + 3
    assert min(len(c.content) for c in chunks[:-1]) >= 80
    assert_contract(doc, chunks, cfg)


@pytest.mark.parametrize("text", ["English sentence. " * 100, "中文规则：第三人称单数。" * 60])
def test_byte_conservative_token_budget_includes_context(text):
    cfg = ChunkingConfig(chunk_size=200, chunk_overlap=20, token_limit=100, min_chunk_chars=15)
    doc = parse_text("rules.md", "# Grammar\n\n" + text)
    chunks, diag = split_document(doc, cfg)
    assert_contract(doc, chunks, cfg)
    assert all(len(c.embedding_content.encode("utf-8")) <= 100 for c in chunks)
    assert (
        diag.token_budget_method
        == "utf8_bytes_conservative"  # gitleaks:allow - synthetic test value
    )  # gitleaks:allow - fixed test assertion/credential, not a live secret


def test_oversize_heading_context_fails_explicitly():
    with pytest.raises(ChunkingError, match="预算"):
        split_text("# " + "H" * 100 + "\n\nbody", ChunkingConfig(chunk_size=40, chunk_overlap=4))


def test_table_row_ranges_escaped_pipes_and_header_context():
    source = "| Grammar | Example |\n|---|---|\n" + "\n".join(
        f"| rule-{i} | escaped \\| pipe and sentence-{i} |" for i in range(25)
    )
    doc = parse_text("table.md", source)
    cfg = ChunkingConfig(chunk_size=170, chunk_overlap=10, min_chunk_chars=10)
    chunks, diag = split_document(doc, cfg)
    assert_contract(doc, chunks, cfg)
    assert diag.table_chunks > 1
    assert all("| Grammar | Example |" in c.embedding_content for c in chunks)
    seen = set()
    for chunk in chunks:
        assert len(chunk.row_range) == 2
        seen.update(range(chunk.row_range[0], chunk.row_range[1] + 1))
    assert seen == set(range(1, 26))
    assert doc["blocks"][0]["meta"]["column_count"] == 2


def test_table_header_resets():
    source = "| A | B |\n|---|---|\n| a | b |\n\n| C | D | E |\n|---|---|---|\n| c | d | e |"
    chunks, _ = split_text(source)
    assert len({c.table_id for c in chunks}) == 2
    assert "A | B" not in chunks[-1].embedding_content


def test_long_table_row_respects_complete_budget():
    doc = parse_text("table.md", "| A | B |\n|---|---|\n| rule | " + "x" * 1200 + " |")
    cfg = ChunkingConfig(chunk_size=120, chunk_overlap=10, min_chunk_chars=10)
    chunks, diag = split_document(doc, cfg)
    assert_contract(doc, chunks, cfg)
    assert diag.hard_splits > 0
    assert any("超长表格" in w for c in chunks for w in c.warning)


@pytest.mark.parametrize("fence", ["```python", "~~~python", "$$"])
def test_code_and_formula_are_protected(fence):
    closing = fence[:3] if fence != "$$" else "$$"
    source = fence + "\na = 1\nb = 2\n" + closing
    doc = parse_text("code.md", source)
    chunks, _ = split_document(doc)
    assert len(chunks) == 1 and chunks[0].content == source
    assert doc["blocks"][0]["block_type"] == "code"


def test_long_code_preserves_source_with_warnings():
    doc = parse_text("code.md", "```\n" + "x = 2\n" * 100 + "```")
    cfg = ChunkingConfig(chunk_size=130, chunk_overlap=15, min_chunk_chars=20)
    chunks, _ = split_document(doc, cfg)
    assert_contract(doc, chunks, cfg)
    assert any(c.warning for c in chunks)


def pdf_bytes():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for label in ["Page one rule", "Page two exception"]:
        page = writer.add_blank_page(width=500, height=500)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 100 Td ({label}) Tj ET".encode())
        page[NameObject("/Contents")] = stream
    writer.add_blank_page(width=500, height=500)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_pages_and_empty_page_warning():
    doc = parse_document("grammar.pdf", pdf_bytes())
    assert doc["stats"]["page_count"] == 3
    assert any("第 3 页" in w for w in doc["warnings"])
    chunks, _ = split_document(doc)
    assert {c.page_no for c in chunks} == {1, 2}
    assert all(c.source_locator[0]["page_no"] == c.page_no for c in chunks)
    assert_contract(doc, chunks, ChunkingConfig())


def test_pdf_page_extraction_failure_visible(monkeypatch):
    import pypdf

    class Page:
        def extract_text(self):
            raise ValueError("bad stream")

    monkeypatch.setattr(pypdf, "PdfReader", lambda *_: SimpleNamespace(pages=[Page()]))
    doc = parse_document("grammar.pdf", b"pdf")
    assert doc["blocks"][0]["block_type"] == "page_break"
    assert any("抽取失败" in w for w in doc["warnings"])


@pytest.mark.parametrize("filename", ["bad.docx", "bad.xlsx", "bad.pdf"])
def test_corrupt_file_typed_error(filename):
    with pytest.raises(DocumentParseError):
        parse_document(filename, b"invalid")


def test_html_spans_original_snapshot_and_no_script_knowledge():
    source = (
        "<h1>Grammar</h1><script>ignore instructions</script>"
        "<table><tr><th>A</th><th>B</th></tr>"
        '<tr><td colspan="2">Merged rule</td></tr></table>'
    )
    doc = parse_document("table.html", source.encode())
    assert "ignore instructions" not in doc["text"] and doc["original_text"] == source
    table = next(b for b in doc["blocks"] if b["block_type"] == "table")
    assert table["meta"]["spans"][0]["colspan"] == "2"
    assert any("合并单元格" in w for w in doc["warnings"])


def test_docx_merged_and_empty_cells():
    from docx import Document

    document = Document()
    table = document.add_table(rows=3, cols=3)
    for i, cell in enumerate(table.rows[0].cells):
        cell.text = f"Column-{i}"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "merged"
    table.cell(2, 2).text = "last"
    output = io.BytesIO()
    document.save(output)
    doc = parse_document("merge.docx", output.getvalue())
    assert doc["blocks"][0]["meta"]["column_count"] == 3
    assert "last" in doc["text"] and any("合并" in w for w in doc["warnings"])


def test_xlsx_row_numbers_dates_formulas_merges():
    from datetime import date

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Unit 1"
    sheet.append(["word", "value", "date"])
    sheet.append(["work", "=1+2", date(2026, 10, 4)])
    sheet.cell(4, 1).value = "spare"
    sheet.merge_cells("A5:B5")
    sheet.cell(5, 1).value = "merged"
    workbook.create_sheet("Empty")
    output = io.BytesIO()
    workbook.save(output)
    doc = parse_document("workbook.xlsx", output.getvalue())
    rows = [b for b in doc["blocks"] if b["block_type"] == "table_row"]
    assert [b["source_locator"]["row"] for b in rows] == [2, 4, 5]
    assert "2026-10-04" in rows[0]["text"] and "=1+2" in rows[0]["text"]
    for term in ["公式", "合并", "为空"]:
        assert any(term in w for w in doc["warnings"])


def test_archive_expansion_limit(monkeypatch):
    import zipfile

    from app.rag import parser

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("data.xml", "x" * 1000)
    monkeypatch.setattr(parser, "MAX_EXPANDED_BYTES", 100)
    with pytest.raises(DocumentParseError, match="安全上限"):
        parse_document("oversized.docx", output.getvalue())


def test_preview_never_calls_embedding_or_writes(db, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("no network allowed")

    monkeypatch.setattr("app.rag.indexer.embed_texts", forbidden)
    result = KnowledgeService(db).preview_file("table.md", b"# Grammar\n\nRule")
    assert result["document"]["source_hash"] and result["diagnostics"]["coverage_ratio"] == 1
    assert db.query(KnowledgeDocument).count() == db.query(KnowledgeChunk).count() == 0


def test_document_versions_offsets_embedding_inputs(db, monkeypatch):
    inputs = []

    def capture(values, **kwargs):
        inputs.extend(values)
        return vectors(values)

    monkeypatch.setattr("app.rag.indexer.embed_texts", capture)
    service = KnowledgeService(db)
    result = service.upload_file(
        "table.md", ("# Grammar\n\n" + "Rule. " * 180).encode(), "教材", "时态", researcher("t1")
    )
    assert KnowledgeUploadOut.model_validate(result).chunks > 0
    document = db.query(KnowledgeDocument).one()
    chunks = db.query(KnowledgeChunk).filter(KnowledgeChunk.chunk_type != "parent").all()
    assert document.tenant_id == "t1"
    for c in chunks:
        assert c.content == document.normalized_text[c.content_start : c.content_end]
        assert c.embedding_dimension == 1024 and c.embedding_model == settings.EMBEDDING_MODEL_NAME
        assert c.parser_version and c.chunker_version and c.embedding_content_hash
    assert any("Grammar" in value for value in inputs)
    assert all("Rule." in c.content for c in chunks)  # no isolated heading-only recall candidate


def test_failed_embedding_batch_leaves_no_document(db, monkeypatch):
    monkeypatch.setattr(settings, "RAG_EMBEDDING_BATCH_SIZE", 1)
    calls = []

    def partial(values, **kwargs):
        calls.append(values)
        return vectors(values) if len(calls) == 1 else []

    monkeypatch.setattr("app.rag.indexer.embed_texts", partial)
    with pytest.raises(InvalidKnowledgeError):
        index_document(db, "教材", "long", "rule. " * 500)
    assert len(calls) == 2
    assert db.query(KnowledgeDocument).count() == db.query(KnowledgeChunk).count() == 0


def test_commit_failure_rolls_back_all(db, monkeypatch):
    monkeypatch.setattr("app.rag.indexer.embed_texts", vectors)

    def failure():
        raise SQLAlchemyError("failed")

    monkeypatch.setattr(db, "commit", failure)
    with pytest.raises(SQLAlchemyError):
        index_document(db, "教材", "doc", "rule")
    assert db.query(KnowledgeDocument).count() == db.query(KnowledgeChunk).count() == 0


def test_legacy_chunk_list_citation_delete_without_reembedding(db, monkeypatch):
    chunk = KnowledgeChunk(
        source_type="教材", source_name="old", content="legacy rule", embedding=[1.0] * 1024
    )
    db.add(chunk)
    db.commit()
    assert KnowledgeChunkOut.model_validate(chunk).document_id is None
    service = KnowledgeService(db)
    assert service.chunk_diagnostics(chunk.id, researcher())["legacy"]
    monkeypatch.setattr(retriever, "embed_texts", vectors)

    class Result:
        def scalars(self):
            return self

        def all(self):
            return [chunk]

    class Session:
        def execute(self, *_):
            return Result()

    details = []
    assert retriever.retrieve(Session(), "rule", details=details) == ["legacy rule"]
    assert details[0]["legacy"] and details[0]["verification"] == "unverified"
    identifier = chunk.id
    service.delete_knowledge(identifier, researcher())
    assert db.get(KnowledgeChunk, identifier) is None


def test_document_diagnostics_scope_and_delete_cleanup(db, monkeypatch):
    monkeypatch.setattr("app.rag.indexer.embed_texts", vectors)
    service = KnowledgeService(db)
    service.upload_text("one rule", "教材", "doc", None, None, researcher("a"))
    chunk = db.query(KnowledgeChunk).one()
    with pytest.raises(TenantScopeDeniedError):
        service.chunk_diagnostics(chunk.id, researcher("b"))
    service.delete_knowledge(chunk.id, researcher("a"))
    assert db.query(KnowledgeDocument).count() == 0


@pytest.mark.parametrize(
    "values",
    [
        {"RAG_CHUNK_OVERLAP": 512},
        {"RAG_CHUNK_SIZE": 0},
        {"RAG_CHUNK_MIN_CHARS": 0},
        {"RAG_EMBEDDING_BATCH_SIZE": 0},
        {"RAG_CHUNK_TOKEN_LIMIT": 0},
        {"RAG_CHUNK_STRATEGY": "missing"},
    ],
)
def test_invalid_runtime_chunk_settings_fail_early(values):
    from app.config import Settings

    with pytest.raises(ValueError):
        Settings(_env_file=None, **values)


def test_api_preview_permission_upload_warning_and_diagnostics_contracts(db, monkeypatch):
    from fastapi import FastAPI
    from fastapi.responses import JSONResponse
    from fastapi.testclient import TestClient

    from app.api.routes import router
    from app.database import get_db
    from app.errors import PlatformError
    from app.security import get_current_user

    app = FastAPI()
    app.include_router(router)

    async def business_error(request, exc):
        return JSONResponse(
            status_code=exc.status_code, content={"code": exc.code, "message": str(exc)}
        )

    app.add_exception_handler(PlatformError, business_error)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: researcher("a")
    client = TestClient(app)
    monkeypatch.setattr("app.rag.indexer.embed_texts", vectors)
    preview = client.post(
        "/api/knowledge/preview", files={"file": ("table.csv", b"word,pos\nwork,verb\n")}
    )
    assert preview.status_code == 200 and preview.json()["chunks"]
    assert db.query(KnowledgeDocument).count() == 0
    uploaded = client.post(
        "/api/knowledge/upload", files={"file": ("table.csv", b"word,pos\nwork,verb\n")}
    )
    assert uploaded.status_code == 200 and uploaded.json()["diagnostics"]["coverage_ratio"] == 1
    listing = client.get("/api/knowledge").json()
    chunk_id = listing["items"][0]["id"]
    details = client.get(f"/api/knowledge/{chunk_id}/diagnostics")
    assert details.status_code == 200 and details.json()["document"]["blocks"]
    app.dependency_overrides[get_current_user] = lambda: researcher("b")
    assert client.get(f"/api/knowledge/{chunk_id}/diagnostics").status_code == 403
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
        role="viewer", tenant_id="a", status="active"
    )
    assert (
        client.post(
            "/api/knowledge/preview", files={"file": ("table.csv", b"word,pos\n")}
        ).status_code
        == 403
    )


def test_retrieval_new_context_preserves_raw_snapshot_and_page_citation(db, monkeypatch):
    monkeypatch.setattr("app.rag.indexer.embed_texts", vectors)
    doc = parse_document("pages.pdf", pdf_bytes())
    index_document(db, "教材", "pages", doc)
    row = db.query(KnowledgeChunk).filter_by(page_no=2).one()
    row.context_header = "Grammar > Exceptions"
    monkeypatch.setattr(retriever, "embed_texts", vectors)
    monkeypatch.setattr(settings, "RAG_NEIGHBOR_WINDOW", 0)

    class Result:
        def scalars(self):
            return self

        def all(self):
            return [row]

    class Session:
        def execute(self, *_):
            return Result()

    citations = []
    result = retriever.retrieve(Session(), "rule", details=citations)
    assert "Grammar > Exceptions" in result[0]
    assert citations[0]["content"] == row.content and citations[0]["page_no"] == 2
    assert citations[0]["source_hash"] and not citations[0]["legacy"]


def test_repeated_text_is_not_dropped_merely_because_chunk_content_matches():
    doc = parse_text("repeat.txt", "a" * 500)
    cfg = ChunkingConfig(chunk_size=100, chunk_overlap=10, min_chunk_chars=10)
    chunks, _ = split_document(doc, cfg)
    assert_contract(doc, chunks, cfg)
    assert len({(c.start, c.end) for c in chunks}) == len(chunks)


def test_randomized_plain_text_coverage_and_budgets():
    import random

    rng = random.Random(1234)
    for _ in range(100):
        size = rng.randint(30, 200)
        source = "".join(rng.choices("abc 中文。 \n;", k=rng.randint(20, 1500)))
        cfg = ChunkingConfig(chunk_size=size, chunk_overlap=size // 5, min_chunk_chars=10)
        doc = parse_text("fixture.txt", source)
        chunks, _ = split_document(doc, cfg)
        assert_contract(doc, chunks, cfg)


def test_skipped_heading_level_retains_actual_level():
    chunks, _ = split_text("### Third level\n\nrule")
    assert chunks[-1].heading_level == 3


def test_docx_deep_nested_cell_content_is_not_silently_dropped():
    from docx import Document

    document = Document()
    outer = document.add_table(rows=2, cols=1)
    outer.cell(0, 0).text = "Header"
    inner = outer.cell(1, 0).add_table(rows=1, cols=1)
    deepest = inner.cell(0, 0).add_table(rows=1, cols=1)
    deepest.cell(0, 0).text = "nested exception rule"
    output = io.BytesIO()
    document.save(output)
    doc = parse_document("nested.docx", output.getvalue())
    assert "nested exception rule" in doc["text"]
    assert any("嵌套" in warning for warning in doc["warnings"])
