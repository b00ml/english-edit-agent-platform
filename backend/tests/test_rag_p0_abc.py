"""P0 RAG A/B/C regression tests: compatibility, parsing, and chunking."""

from __future__ import annotations

import io

import pytest

from app.rag.chunking import ChunkingConfig, split_text
from app.rag.parser import parse_document, parse_file


def _docx_bytes() -> bytes:
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("段落一")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "知识点"
    table.cell(0, 1).text = "例句|说明"
    table.cell(1, 0).text = "一般现在时"
    table.cell(1, 1).text = "I work"
    table.cell(2, 0).text = "第三人称"
    table.cell(2, 1).text = "He works"
    document.add_paragraph("段落二")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_parser_keeps_body_order_and_table_content():
    parsed = parse_document("lesson.docx", _docx_bytes())
    assert [block["block_type"] for block in parsed["blocks"]] == [
        "paragraph",
        "table",
        "paragraph",
    ]
    table = parsed["blocks"][1]
    assert "一般现在时" in table["text"]
    assert "例句\\|说明" in table["text"]
    assert table["meta"]["column_count"] == 2
    assert "一般现在时" in parse_file("lesson.docx", _docx_bytes())


def test_markdown_table_is_normalized_and_keeps_all_rows():
    source = "# 时态\n\n| 项目 | 用法 |\n|---|---|\n" + "\n".join(
        f"| rule-{index} | example-{index} |" for index in range(12)
    )
    parsed = parse_document("grammar.md", source.encode())
    assert len(parsed["blocks"]) == 2
    table = parsed["blocks"][1]
    assert table["meta"]["row_count"] == 12
    assert "| 项目 | 用法 |" in table["text"]
    assert "| --- | --- |" in table["text"]
    assert all(f"rule-{index}" in table["text"] for index in range(12))


def test_html_table_is_converted_to_gfm():
    html = (
        "<h1>时态</h1><p>说明</p>"
        "<table><tr><th>主语</th><th>动词</th></tr>"
        "<tr><td>He</td><td>works</td></tr></table>"
    )
    parsed = parse_document("grammar.html", html.encode())
    tables = [block for block in parsed["blocks"] if block["block_type"] == "table"]
    assert len(tables) == 1
    assert "| 主语 | 动词 |" in tables[0]["text"]
    assert "| He | works |" in tables[0]["text"]


def test_csv_parser_preserves_sheet_columns_and_row_numbers():
    parsed = parse_document("vocab.csv", "word,pos\nwork,verb\n".encode())
    assert parsed["blocks"][0]["meta"]["chunk_type"] == "table_summary"
    row = next(b for b in parsed["blocks"] if b["block_type"] == "table_row")
    assert "行号：2" in row["text"]
    assert "word：work" in row["text"]
    assert "pos：verb" in row["text"]
    assert row["source_locator"]["row"] == 2


def test_xlsx_parser_reads_real_workbook_and_locations():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Unit 1"
    sheet.append(["word", "pos"])
    sheet.append(["work", "verb"])
    output = io.BytesIO()
    workbook.save(output)
    parsed = parse_document("vocab.xlsx", output.getvalue())
    assert any("Sheet：Unit 1" in b["text"] for b in parsed["blocks"])
    assert any("word：work" in b["text"] for b in parsed["blocks"])


def test_chunker_fixes_newline_pathology_and_respects_hard_limit():
    text = "A" * 480 + "\n" + "B" * 900
    chunks, diagnostics = split_text(
        text,
        ChunkingConfig(chunk_size=100, chunk_overlap=10, min_chunk_chars=20),
    )
    assert chunks
    assert all(len(chunk.content) <= 100 for chunk in chunks)
    assert min(len(chunk.content) for chunk in chunks[:-1]) >= 20
    assert diagnostics.hard_splits > 0


def test_chunker_preserves_heading_context_and_table_headers():
    text = "# 英语语法\n\n## 一般现在时\n\n规则说明。\n\n| 项目 | 例句 |\n|---|---|\n" + "\n".join(
        f"| rule-{index} | example-{index} |" for index in range(20)
    )
    chunks, diagnostics = split_text(
        text,
        ChunkingConfig(chunk_size=110, chunk_overlap=10, min_chunk_chars=10),
    )
    assert diagnostics.table_chunks > 0
    table_chunks = [chunk for chunk in chunks if chunk.table_id]
    assert table_chunks
    assert all("| 项目 | 例句 |" in chunk.embedding_content for chunk in table_chunks)
    assert all(chunk.section_path for chunk in chunks)
    assert any(chunk.section_path == ["英语语法", "一般现在时"] for chunk in chunks)


def test_chunker_rejects_invalid_overlap():
    with pytest.raises(ValueError, match="chunk_overlap"):
        split_text("content", ChunkingConfig(chunk_size=10, chunk_overlap=10))
